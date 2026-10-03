import openpyxl
from datetime import datetime, date
from pathlib import Path
from typing import Dict, Any, List, Optional
from app.config import settings

class ExcelWriter:
    def __init__(self, excel_path: Optional[str] = None):
        self.excel_path = Path(excel_path) if excel_path else settings.get_excel_absolute_path()

    def sync_detected_events(self, detected_events: List[Dict[str, Any]]) -> Dict[str, Any]:
        """
        Synchronize detected event results into bybit-okx_maliyet.xlsx.
        Updates existing rows or appends to the first empty row.
        """
        if not self.excel_path.exists():
            return {"success": False, "error": f"Excel dosyası bulunamadı: {self.excel_path}"}

        try:
            wb = openpyxl.load_workbook(self.excel_path)
        except Exception as e:
            return {"success": False, "error": f"Excel dosyası açılamadı (açıksa kapatınız): {e}"}

        try:
            updated_count = 0
            added_count = 0

            for item in detected_events:
                exchange = item["exchange"]
                sheet_name = "Bybit Maliyetleri" if str(exchange).lower() == "bybit" else "OKX Maliyetleri"
                if sheet_name not in wb.sheetnames:
                    continue

                ws = wb[sheet_name]
                coin = item["coin"].upper()
                event_type = item["event_type"]
                market = item["market"]
                vol = float(item["total_volume"])
                cost = float(item["total_cost"])
                fee_rate = float(item.get("fee_rate", 0.0))
                
                # Format date
                date_str = item.get("date", datetime.now().strftime("%Y-%m-%d"))
                try:
                    dt_obj = datetime.strptime(date_str, "%Y-%m-%d").date()
                except Exception:
                    dt_obj = date.today()

                # Find if this event already exists in rows 12..5000 (match coin, type, market AND date)
                row_to_write = None
                first_empty_row = None

                for r in range(12, 5001):
                    cell_coin = str(ws[f"D{r}"].value or "").strip().upper()
                    cell_event = str(ws[f"B{r}"].value or "").strip()
                    cell_market = str(ws[f"C{r}"].value or "").strip()
                    raw_d = ws[f"A{r}"].value
                    cell_date_str = ""
                    if isinstance(raw_d, (datetime, date)):
                        cell_date_str = raw_d.strftime("%Y-%m-%d")
                    elif raw_d:
                        cell_date_str = str(raw_d).strip()[:10]

                    if not cell_coin and first_empty_row is None:
                        first_empty_row = r

                    # Match existing event on coin, event_type, market and date
                    if (cell_coin == coin and 
                        (cell_event == event_type or not cell_event) and 
                        (cell_market == market or not cell_market) and
                        (cell_date_str == date_str or not cell_date_str)):
                        row_to_write = r
                        break

                if row_to_write is not None:
                    # Update existing row
                    ws[f"E{row_to_write}"] = vol
                    ws[f"F{row_to_write}"] = cost
                    if fee_rate > 0:
                        ws[f"J{row_to_write}"] = fee_rate
                    updated_count += 1
                elif first_empty_row is not None:
                    # Append to first empty row
                    r = first_empty_row
                    ws[f"A{r}"] = dt_obj
                    ws[f"B{r}"] = event_type
                    ws[f"C{r}"] = market
                    ws[f"D{r}"] = coin
                    ws[f"E{r}"] = vol
                    ws[f"F{r}"] = cost
                    if fee_rate > 0:
                        ws[f"J{r}"] = fee_rate

                    # Ensure formulas are always present on any appended row
                    ws[f"G{r}"] = f'=IFERROR(IF(E{r}>0,(F{r}/E{r})*1000,""),"")'
                    if str(exchange).lower() == "bybit":
                        ws[f"H{r}"] = f'=IFERROR(IF(E{r}>0,E{r}*IF(ISNUMBER(J{r}),J{r},IF(OR(C{r}="Alpha",B{r}="Alpha"),$C$7,IF(C{r}="Futures",$C$6,$C$5))),""),"")'
                        ws[f"K{r}"] = f'=IFERROR(IF(H{r}>0,H{r}*IF(OR(C{r}="Alpha",B{r}="Alpha"),$C$9,$C$8),""),"")'
                    else:
                        ws[f"H{r}"] = f'=IFERROR(IF(E{r}>0,E{r}*IF(ISNUMBER(J{r}),J{r},IF(C{r}="Futures",$C$6,$C$5)),""),"")'
                        ws[f"K{r}"] = f'=IFERROR(IF(H{r}>0,H{r}*$C$7,""),"")'
                    ws[f"I{r}"] = f'=IFERROR(IF(AND(ISNUMBER(F{r}),ISNUMBER(H{r})),F{r}-H{r},""),"")'
                    ws[f"L{r}"] = f'=IFERROR(IF(ISNUMBER(F{r}),F{r}-IF(ISNUMBER(K{r}),K{r},0),""),"")'
                    added_count += 1

            try:
                wb.save(self.excel_path)
                return {
                    "success": True,
                    "added": added_count,
                    "updated": updated_count,
                    "file": str(self.excel_path)
                }
            except PermissionError:
                return {
                    "success": False,
                    "error": "Excel dosyası başka bir programda (Excel) açık! Lütfen dosyayı kaydedip kapatıp tekrar senkronize ediniz."
                }
            except Exception as e:
                return {"success": False, "error": str(e)}
        finally:
            wb.close()
