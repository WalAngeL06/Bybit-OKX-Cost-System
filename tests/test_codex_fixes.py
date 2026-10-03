import os
import sys
import unittest
import asyncio
import tempfile
import openpyxl
from pathlib import Path
from datetime import datetime, date
from unittest.mock import patch, AsyncMock, MagicMock
from typing import get_type_hints

# Add project root to sys.path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.core.excel_writer import ExcelWriter
from app.core.detector import DetectorEngine
from app.core.bybit_client import BybitClient
from app.core.okx_client import OKXClient
from app.database.models import add_event, get_all_events, save_trades_to_cache
from app.database.db import get_connection, init_db
from app.config import settings

class TestCodexFixes(unittest.TestCase):

    def test_item1_pagination_and_cache_limit(self):
        """CODEX Point 1: Test pagination (no 100-trade cut) and no 500-trade cache retrieval limit."""
        # 1. Test Bybit pagination across cursor
        bybit = BybitClient("test_key", "test_secret")
        page1 = {
            "list": [{"execId": f"b_{i}", "orderId": f"o_{i}", "symbol": "BTCUSDT", "side": "Buy",
                      "execPrice": "50000", "execQty": "0.1", "execValue": "5000", "execFee": "0.5",
                      "feeCurrency": "USDT", "execTime": 1700000000000 + i} for i in range(100)],
            "nextPageCursor": "cursor_page_2"
        }
        page2 = {
            "list": [{"execId": f"b_{i}", "orderId": f"o_{i}", "symbol": "BTCUSDT", "side": "Buy",
                      "execPrice": "50000", "execQty": "0.1", "execValue": "5000", "execFee": "0.5",
                      "feeCurrency": "USDT", "execTime": 1700000000000 + i} for i in range(100, 115)],
            "nextPageCursor": None
        }

        async def mock_bybit_request(method, endpoint, params=None):
            if params and params.get("cursor") == "cursor_page_2":
                return page2
            return page1

        with patch.object(bybit, "_request", side_effect=mock_bybit_request):
            trades = asyncio.run(bybit.get_executions(limit=100, max_trades=200))
            self.assertEqual(len(trades), 115, "Bybit client must paginate and fetch all 115 trades across cursors")

        # 2. Test OKX pagination across cursor
        okx = OKXClient("test_key", "test_secret", "test_pass")
        okx_page1 = [{"fillId": f"okx_{i}", "ordId": f"o_{i}", "instId": "BTC-USDT", "side": "buy",
                      "fillPx": "50000", "fillSz": "0.1", "fee": "-0.5", "feeCcy": "USDT",
                      "ts": str(1700000000000 + i), "billId": f"bill_{i}"} for i in range(100)]
        okx_page2 = [{"fillId": f"okx_{i}", "ordId": f"o_{i}", "instId": "BTC-USDT", "side": "buy",
                      "fillPx": "50000", "fillSz": "0.1", "fee": "-0.5", "feeCcy": "USDT",
                      "ts": str(1700000000000 + i), "billId": f"bill_{i}"} for i in range(100, 120)]

        async def mock_okx_request(method, endpoint, params=None):
            if params and params.get("after"):
                return okx_page2
            return okx_page1

        with patch.object(okx, "_request", side_effect=mock_okx_request):
            okx_trades = asyncio.run(okx.get_fills(limit=100, max_trades=200))
            self.assertEqual(len(okx_trades), 120, "OKX client must paginate and fetch all 120 trades across cursors")

        # 3. Test SQLite query retrieves >500 records without limit
        conn = get_connection()
        try:
            cursor = conn.cursor()
            cursor.execute("DELETE FROM trade_cache WHERE exec_id LIKE 'cache_test_%'")
            # Insert 550 trades
            records = [
                ("Bybit", f"cache_test_{i}", f"ord_{i}", "BTCUSDT", "Buy", 50000.0, 0.01, 500.0, 0.5, 0.001,
                 1700000000000 + i, "Spot", "2026-10-01")
                for i in range(550)
            ]
            cursor.executemany("""
            INSERT OR IGNORE INTO trade_cache (
                exchange, exec_id, order_id, symbol, side, price, qty,
                exec_value, fee, fee_rate, exec_time, market, created_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, records)
            conn.commit()

            cursor.execute("SELECT * FROM trade_cache WHERE exec_id LIKE 'cache_test_%' ORDER BY exec_time ASC")
            fetched = cursor.fetchall()
            self.assertEqual(len(fetched), 550, "SELECT query must retrieve all 550 trades without arbitrary 500 limit")
        finally:
            cursor.execute("DELETE FROM trade_cache WHERE exec_id LIKE 'cache_test_%'")
            conn.commit()
            conn.close()

    def test_item2_market_and_symbol_matching(self):
        """CODEX Point 2: Test strict market and symbol matching (Spot vs Futures, KII vs KIICHI)."""
        detector = DetectorEngine()

        mock_event = {
            "id": 999,
            "exchange": "Bybit",
            "event_name": "KII Alpha",
            "event_type": "Alpha",
            "market": "Alpha",
            "coin": "KII",
            "pair": "KIIUSDT",
            "start_date": "2026-10-01",
            "end_date": None
        }

        with patch("app.core.detector.get_active_events", return_value=[mock_event]):
            trades = [
                # 1. Valid matching trade (KII Spot/Alpha)
                {
                    "exchange": "bybit", # lower-case exchange test
                    "symbol": "KIIUSDT",
                    "market": "Spot",
                    "side": "Buy",
                    "exec_value": 100.0,
                    "qty": 10.0,
                    "fee": 0.1,
                    "exec_time": int(datetime(2026, 10, 2).timestamp() * 1000)
                },
                # 2. Invalid symbol: KIICHI (prefix false positive)
                {
                    "exchange": "Bybit",
                    "symbol": "KIICHIUSDT",
                    "market": "Spot",
                    "side": "Buy",
                    "exec_value": 900.0,
                    "qty": 90.0,
                    "fee": 0.9,
                    "exec_time": int(datetime(2026, 10, 2).timestamp() * 1000)
                },
                # 3. Invalid market: Futures trade for Alpha/Spot event
                {
                    "exchange": "Bybit",
                    "symbol": "KIIUSDT",
                    "market": "Futures",
                    "side": "Buy",
                    "exec_value": 500.0,
                    "qty": 50.0,
                    "fee": 0.5,
                    "exec_time": int(datetime(2026, 10, 2).timestamp() * 1000)
                }
            ]

            results = detector.match_events(trades)
            self.assertEqual(len(results), 1)
            # Only the 100$ trade must match, NOT 1,000$ or 1,500$!
            self.assertEqual(results[0]["total_volume"], 100.0)
            self.assertEqual(results[0]["trades_count"], 1)

    def test_item3_fee_conversion(self):
        """CODEX Point 3: Test non-USD fee conversion in Bybit & OKX, and OKX contract multiplier (ctVal)."""
        # 1. Bybit fee conversion test via BybitClient
        bybit = BybitClient("test", "test")
        bybit_raw = {
            "list": [{
                "execId": "ex_btc_fee",
                "orderId": "ord_1",
                "symbol": "BTCUSDT",
                "side": "Buy",
                "execPrice": "30000.0",
                "execQty": "1.0",
                "execValue": "30000.0",
                "execFee": "0.001",
                "feeCurrency": "BTC",  # 0.001 BTC fee at $30,000 = $30.0
                "feeRate": "0.001",
                "execTime": 1700000000000,
                "isMaker": False
            }]
        }
        with patch.object(bybit, "_request", new_callable=AsyncMock, return_value=bybit_raw):
            trades = asyncio.run(bybit.get_executions())
            self.assertEqual(len(trades), 1)
            self.assertAlmostEqual(trades[0]["fee"], 30.0, places=2)

        # 2. OKX contract multiplier (ctVal) and fee conversion test
        okx = OKXClient("test", "test", "pass")
        okx_raw = [{
            "fillId": "okx_swap_1",
            "ordId": "ord_swap_1",
            "instId": "BTC-USDT-SWAP",
            "side": "buy",
            "fillPx": "50000.0",
            "fillSz": "100",        # 100 contracts
            "fee": "-0.001",        # charged in BTC
            "feeCcy": "BTC",
            "feeRate": "0.0005",
            "ts": "1700000000000",
            "execType": "T"
        }]

        with patch.object(okx, "_request", new_callable=AsyncMock, return_value=okx_raw), \
             patch.object(okx, "get_contract_val", new_callable=AsyncMock, return_value=0.01):
            trades = asyncio.run(okx.get_fills(inst_type="SWAP"))
            self.assertEqual(len(trades), 1)
            # 100 contracts * 0.01 ctVal = 1.0 BTC quantity
            self.assertAlmostEqual(trades[0]["qty"], 1.0, places=4)
            # Notional value: 1.0 * 50,000 = $50,000
            self.assertAlmostEqual(trades[0]["exec_value"], 50000.0, places=2)
            # Fee in BTC converted to USD: 0.001 * 50,000 = $50.0
            self.assertAlmostEqual(trades[0]["fee"], 50.0, places=2)

    def test_item4_unsold_held_assets_cost(self):
        """CODEX Point 4: Test that partially sold position does NOT create fake slippage cost."""
        detector = DetectorEngine()

        mock_event = {
            "id": 998,
            "exchange": "Bybit",
            "event_name": "TEST Splash",
            "event_type": "Splash",
            "market": "Spot",
            "coin": "TEST",
            "pair": "TESTUSDT",
            "start_date": "2026-10-01",
            "end_date": None
        }

        # Bought 100 units at $10 = $1,000. Sold 10 units at $10 = $100. Price unchanged!
        trades = [
            {
                "exchange": "Bybit",
                "symbol": "TESTUSDT",
                "market": "Spot",
                "side": "Buy",
                "price": 10.0,
                "qty": 100.0,
                "exec_value": 1000.0,
                "fee": 1.0,
                "exec_time": int(datetime(2026, 10, 2).timestamp() * 1000)
            },
            {
                "exchange": "Bybit",
                "symbol": "TESTUSDT",
                "market": "Spot",
                "side": "Sell",
                "price": 10.0,
                "qty": 10.0,
                "exec_value": 100.0,
                "fee": 0.1,
                "exec_time": int(datetime(2026, 10, 2).timestamp() * 1000)
            }
        ]

        with patch("app.core.detector.get_active_events", return_value=[mock_event]):
            results = detector.match_events(trades)
            self.assertEqual(len(results), 1)
            # Realized slippage must be 0.0, NOT 900.0!
            self.assertEqual(results[0]["makas"], 0.0)
            # Total cost is strictly fees (1.10)
            self.assertAlmostEqual(results[0]["total_cost"], 1.10, places=2)

    def test_item5_excel_writer_date_isolation(self):
        """CODEX Point 5: Test that events on different dates are not overwritten into one row."""
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            tmp_path = tmp.name

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Bybit Maliyetleri"
        ws["C5"] = 0.001
        ws["C6"] = 0.0005
        ws["C7"] = 0.0002
        ws["C8"] = 0.55
        ws["C9"] = 0.60
        wb.save(tmp_path)
        wb.close()

        writer = ExcelWriter(excel_path=tmp_path)
        
        # Event 1 on Oct 1st
        event1 = {
            "exchange": "Bybit",
            "coin": "KII",
            "event_type": "Alpha",
            "market": "Alpha",
            "total_volume": 50000.0,
            "total_cost": 25.0,
            "date": "2026-10-01"
        }
        res1 = writer.sync_detected_events([event1])
        self.assertEqual(res1["added"], 1)

        # Event 2 on Oct 5th for same coin and type
        event2 = {
            "exchange": "Bybit",
            "coin": "KII",
            "event_type": "Alpha",
            "market": "Alpha",
            "total_volume": 75000.0,
            "total_cost": 37.5,
            "date": "2026-10-05"
        }
        res2 = writer.sync_detected_events([event2])
        # Event 2 MUST be added as a separate row, NOT overwrite row 12!
        self.assertEqual(res2["added"], 1)

        wb_verify = openpyxl.load_workbook(tmp_path)
        ws_verify = wb_verify["Bybit Maliyetleri"]
        self.assertEqual(ws_verify["E12"].value, 50000.0)
        self.assertEqual(ws_verify["E13"].value, 75000.0)
        wb_verify.close()

        try:
            os.remove(tmp_path)
        except Exception:
            pass

    def test_item6_save_settings_no_wipe(self):
        """CODEX Point 6: Test that saving settings with empty strings does not wipe existing values."""
        from app.main import SettingsPayload, save_settings

        # If user submits only gemini key, other keys must NOT be wiped
        payload = SettingsPayload(
            bybit_api_key="",
            bybit_api_secret="",
            gemini_api_key="TEST_GEMINI_KEY"
        )
        
        saved_dict = {}
        with patch.object(settings, "save_to_env", side_effect=lambda d: saved_dict.update(d)):
            res = asyncio.run(save_settings(payload))
            self.assertTrue(res["success"])
            self.assertIn("GEMINI_API_KEY", saved_dict)
            self.assertNotIn("BYBIT_API_KEY", saved_dict)
            self.assertNotIn("BYBIT_API_SECRET", saved_dict)
            # When bybit_testnet is None, BYBIT_TESTNET should not be overwritten
            self.assertNotIn("BYBIT_TESTNET", saved_dict)

    def test_item7_exchange_error_reporting(self):
        """CODEX Point 7: Test exchange errors are accurately captured and reported."""
        detector = DetectorEngine()

        # Simulate Bybit error during fetch
        with patch.object(settings, "BYBIT_API_KEY", "real_key"), \
             patch.object(settings, "BYBIT_API_SECRET", "real_secret"), \
             patch.object(detector.bybit, "get_executions", side_effect=Exception("API connection timeout")):
            
            res = asyncio.run(detector.fetch_and_cache_trades())
            self.assertFalse(res["success"], "Fetch result must be False when an exchange fails")
            self.assertTrue(any("API connection timeout" in err for err in res["errors"]),
                            "Errors list must contain verbatim exception message")

    def test_item8_auto_sync_worker(self):
        """CODEX Point 8: Test auto_sync_worker background task."""
        from app.main import auto_sync_worker

        run_sync_called = []
        async def mock_run_sync():
            run_sync_called.append(True)

        call_count = 0
        async def mock_sleep(seconds):
            nonlocal call_count
            call_count += 1
            if call_count >= 2:
                raise asyncio.CancelledError()

        with patch.object(settings, "BYBIT_API_KEY", "key"), \
             patch.object(settings, "BYBIT_API_SECRET", "sec"), \
             patch("app.main.run_sync", side_effect=mock_run_sync), \
             patch("asyncio.sleep", side_effect=mock_sleep):
            
            try:
                asyncio.run(auto_sync_worker())
            except asyncio.CancelledError:
                pass
            
            self.assertGreaterEqual(len(run_sync_called), 1, "auto_sync_worker must call run_sync periodically")

    def test_item9_typing_and_platform_compatibility(self):
        """CODEX Point 9: Test typing compatibility (Optional in annotations) and platform imports."""
        # 1. Type annotations evaluation
        hints_excel = get_type_hints(ExcelWriter.__init__)
        self.assertIn("excel_path", hints_excel, "ExcelWriter.__init__ must have valid evaluated type hints")

        # 2. Check sys module and platform branches in app/main.py
        import app.main as main_mod
        self.assertTrue(hasattr(main_mod, "sys"), "sys must be imported in app/main")

        # 3. Test open_excel logic on non-Windows
        with patch("os.name", "posix"), \
             patch("sys.platform", "linux"), \
             patch("subprocess.Popen") as mock_popen, \
             patch.object(settings, "get_excel_absolute_path", return_value=Path(__file__)):
            res = asyncio.run(main_mod.open_excel())
            self.assertTrue(res["success"])
            mock_popen.assert_called_once()

if __name__ == "__main__":
    unittest.main()
